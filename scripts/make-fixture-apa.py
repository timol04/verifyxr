"""Erzeugt fixtures/apa.docx: Fliesstext mit allen Zitationsformen plus Literaturverzeichnis
mit bewusst eingebauten Fehlern (siehe test/apa.test.mjs für die erwarteten Befunde)."""
from docx import Document

doc = Document()
doc.add_heading('Einleitung', level=1)

def para(parts):
    """parts: Liste aus str oder (str, 'i') für kursiv."""
    p = doc.add_paragraph()
    for part in parts:
        text, style = (part, '') if isinstance(part, str) else part
        r = p.add_run(text)
        r.italic = style == 'i'
    return p

para(['Gruppenlernen wirkt (Meier, 2020a). Zusammenarbeit hilft ebenfalls (Meier & Huber, 2021), und Stress wird bewältigt (Müller et al., 2019).'])
para(['Mehrere Quellen in einer Klammer (Meier, 2020a; Keller, 2018) stützen das. Meier und Huber (2021) zeigen zudem Effekte.'])
para(['Fehlerfälle: Es gibt Hinweise (Schulz, 2017), Keller (2017) widerspricht, und (Meier, 2020b) fehlt ebenfalls.'])
para(['Ohne Zitation: Das Wetter war schön (siehe Abbildung 1) und die Fallzahl betrug (n = 20).'])

doc.add_heading('Literaturverzeichnis', level=1)
para(['Keller, E. (2018). Das Gehirn und seine Wunder. Neuro-Verlag.'])
para(['Huber, B. (2022). Unused Reference On Many Topics. ', ('Journal of Things', 'i'), ', ', ('5', 'i'), '(1), 1–10. https://doi.org/10.2222/zzzz'])
para(['Meier, A. (2020a). Learning in groups. ', ('Journal of Education', 'i'), ', ', ('12', 'i'), '(3), 45–67. https://doi.org/10.1234/abcd'])
para(['Meier, A., & Huber, B. (2021). Cooperation at school. ', ('Educational Psychology Review', 'i'), ', ', ('33', 'i'), ', 1–20. https://doi.org/10.5678/efgh'])
para(['Müller, C., Schmid, D., & Keller, E. (2019). Stress and coping. Psychology Today, 10(2), 5–9. doi:10.1111/xyz'])

doc.save('fixtures/apa.docx')
print('fixtures/apa.docx geschrieben')
