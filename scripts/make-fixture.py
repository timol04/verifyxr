"""Erzeugt fixtures/beispiel.docx mit bewusst unterschiedlichen Formatierungs-Quellen,
damit das Inspect-Script die OOXML-Vererbung (docDefaults -> Stil -> basedOn -> direkt) zeigen kann."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

doc = Document()
styles = doc.styles

# 1) docDefaults: Basis für alles, was nirgends sonst gesetzt ist
rpr_def = styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
rfonts = rpr_def.find(qn('w:rFonts'))
for a in list(rfonts.attrib):
    del rfonts.attrib[a]
rfonts.set(qn('w:ascii'), 'Arial'); rfonts.set(qn('w:hAnsi'), 'Arial')
rpr_def.find(qn('w:sz')).set(qn('w:val'), '22')  # 11pt

# 2) Normal: Times New Roman 12pt, Zeilenabstand 1.5 (line=360, auto)
normal = styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5

# 3) Eigener Stil "Fliesstext" basedOn Normal, überschreibt nur den Zeilenabstand (doppelt)
body = styles.add_style('Fliesstext', 1)
body.base_style = normal
body.paragraph_format.line_spacing = 2.0

# Heading 1 erbt von Normal, setzt aber eigene Grösse
styles['Heading 1'].font.size = Pt(16)
styles['Heading 1'].font.name = 'Times New Roman'

# Seite 1: Ränder 2.5 cm
s = doc.sections[0]
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)
s.header_distance = s.footer_distance = Cm(1.25)

doc.add_heading('Einleitung', level=1)
doc.add_paragraph('Absatz im Stil Normal: erbt alles aus dem Normal-Stil (TNR 12pt, 1.5-fach).')
doc.add_paragraph('Absatz im Stil Fliesstext: basedOn Normal, aber doppelter Zeilenabstand.', style='Fliesstext')

# Direkte Formatierung: Absatz-Spacing 1.0 + Run mit Courier 10pt
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.0
r = p.add_run('Absatz mit direkter Formatierung: Courier New 10pt, einfacher Zeilenabstand.')
r.font.name = 'Courier New'; r.font.size = Pt(10)

# Mehrere Runs mit gemischter Formatierung im selben Absatz
p = doc.add_paragraph('Gemischt: normal, ')
r = p.add_run('gross (14pt)'); r.font.size = Pt(14)
p.add_run(' und wieder normal.')

# Zeilenabstand "mindestens/genau" in Punkt (lineRule=exact, 18pt = 360)
p = doc.add_paragraph('Absatz mit exaktem Zeilenabstand von 18pt (lineRule=exact).')
p.paragraph_format.line_spacing = Pt(18)  # python-docx setzt dann lineRule=exact

# Abschnittswechsel: Abschnitt 2 mit anderen Rändern (3 cm links/rechts)
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
sec2.left_margin = sec2.right_margin = Cm(3)
sec2.top_margin = sec2.bottom_margin = Cm(2.5)
doc.add_paragraph('Abschnitt 2 (neue Seite): andere Seitenränder, Stil Normal.')

doc.save('fixtures/beispiel.docx')
print('fixtures/beispiel.docx geschrieben')
